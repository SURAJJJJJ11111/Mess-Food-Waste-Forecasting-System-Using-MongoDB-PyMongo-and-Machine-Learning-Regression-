"""
Academic Presentation Generator
Creates a professional 12-slide PowerPoint presentation (.pptx)
for the CSE494 CA2 project evaluation and technical viva, including
the dual machine learning algorithm comparison (Random Forest vs Decision Tree).
"""

import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import config


def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette
    PRIMARY_COLOR = RGBColor(30, 60, 114)    # Deep Navy
    ACCENT_COLOR = RGBColor(139, 92, 246)    # Violet / Purple
    TEXT_COLOR = RGBColor(31, 41, 55)        # Charcoal
    MUTED_COLOR = RGBColor(107, 114, 128)    # Gray
    CARD_BG_COLOR = RGBColor(248, 250, 252)  # Soft gray

    slides_content = [
        {
            "title": "Mess Food Waste Forecasting System",
            "subtitle": "Using MongoDB, PyMongo, and Comparative Machine Learning Regression\n(Random Forest Regressor vs Decision Tree Regressor)\nCSE494: Intelligent NoSQL Databases — CA2 Assessment | Lovely Professional University",
            "bullets": [
                "Individual Academic Project Submission for CA2 & Technical Viva Examination",
                "Technology Stack: MongoDB Server 8.x, PyMongo v4, scikit-learn, Streamlit, Pandas, Plotly",
                "Dual Regression ML Benchmark: Random Forest (14.229 kg MAE) vs Decision Tree (14.645 kg MAE)",
                "Full-Stack Operational Solution: Strict Schema Validation, Native Aggregations, & Real-Time Dashboard"
            ]
        },
        {
            "title": "Problem Statement & Motivation",
            "subtitle": "Operational & Economic Uncertainty in Large-Scale Student Dining",
            "bullets": [
                "The Challenge: University hostel catering serves 800+ diners per session with volatile daily attendance.",
                "Economic Loss: Kitchen over-preparation generates tons of edible food waste (₹45–₹65/kg operational cost).",
                "Shortage Risk: Under-preparation leads to student dissatisfaction and operational emergencies.",
                "Technical Opportunity: Exploit NoSQL document flexibility to store complex meal sessions and comparative ML to forecast surplus food prior to cooking."
            ]
        },
        {
            "title": "Project Objectives & Technical Scope",
            "subtitle": "Core Deliverables for Intelligent NoSQL Implementation",
            "bullets": [
                "1. Database Architecture: Implement 4 MongoDB collections with embedded structures and performance indexes.",
                "2. Dynamic Aggregation: Design 6 analytical pipelines for meal types, daily trends, itemized unwinding, and financial loss.",
                "3. Dual Regression ML Benchmark: Implement Random Forest and Decision Tree regressors on identical test sets.",
                "4. High-Waste Risk Classification: Build an auxiliary classifier (>35 kg threshold) to trigger proactive kitchen warnings.",
                "5. Operational Dashboard: Deliver a multi-page interactive Streamlit dashboard for real-time forecasting and CRUD control."
            ]
        },
        {
            "title": "System Architecture & Data Pipeline",
            "subtitle": "Three-Tier Decoupled Architecture",
            "bullets": [
                "Data Tier: MongoDB Community Server holding menus, meal_records, attendance_records, predictions.",
                "Validation Layer: Native $jsonSchema rules enforcing non-negative weights, enum meal types, and required fields.",
                "ML & Analytics Tier: Scikit-learn Pipeline with ColumnTransformer, StandardScaler, and dual regression models.",
                "Application Tier: Streamlit multi-page interface connecting via PyMongo driver over TCP with interactive Plotly charts."
            ]
        },
        {
            "title": "MongoDB Schema Design & Validation",
            "subtitle": "Document Modeling Principles & Constraint Enforcement",
            "bullets": [
                "Embedded Documents: 1:Few relationship; food_items array embedded inside meal_records for atomic reads without multi-table joins.",
                "Referenced Documents: attendance_records and predictions collections referenced by ID to maintain separation of concerns.",
                "Server-Side $jsonSchema: Enforces required attributes, strict BSON types, and non-negative bounds at database engine level.",
                "Unique & Compound Indexes: idx_unique_meal_id and compound (meal_date, meal_type) prevent duplicate records and speed queries."
            ]
        },
        {
            "title": "PyMongo CRUD & Aggregation Demonstrations",
            "subtitle": "Analytical Intelligence Extracted Directly in MongoDB",
            "bullets": [
                "Pipeline 1 (Meal Types): Lunch produces highest average leftover (58.25 kg / 13.6%), followed by Dinner (50.06 kg / 13.2%).",
                "Pipeline 2 (Daily Trends): Groups by calendar date to track daily prepared vs leftover quantities and waste percentages.",
                "Pipeline 3 (Item Unwind): Uses $unwind on food_items to identify high-waste menu staples (rice and curries).",
                "Pipeline 4 (Attendance Bands): Reveals that low-attendance sessions (<450 diners) suffer from highest relative over-preparation.",
                "Pipeline 6 (Financial Loss): Projects monthly cost of avoidable waste (~₹1.9L - ₹2.24L INR/month @ ₹45/kg)."
            ]
        },
        {
            "title": "Machine Learning Workflow & Preprocessing",
            "subtitle": "Zero-Leakage Training Strategy on 1,500 Historical Meals",
            "bullets": [
                "Feature Selection: Strictly uses pre-meal features (expected attendance, planned kg, meal type, day of week, weather, event flag, lag leftovers).",
                "Leakage Prevention: Actual attendance, served kg, and post-meal waste are strictly excluded from prediction features.",
                "Chronological Splitting: 80% older meals for training (1,199 records), 20% newest meals for testing (301 records).",
                "Physical Output Clipping: Predictions constrained to >= 0 kg to eliminate physically impossible negative values."
            ]
        },
        {
            "title": "ML Algorithm 1: Random Forest Regressor",
            "subtitle": "Ensemble Bagging Architecture for Surplus Food Prediction",
            "bullets": [
                "Methodology: Ensembles 100 de-correlated decision trees using bootstrap aggregation (bagging) and random feature subsets.",
                "Test Performance: MAE = 14.229 kg | RMSE = 19.514 kg | R² Score = 0.4959.",
                "Baseline Gain: Outperforms historical meal-type baseline (15.412 kg MAE) by +7.68% error reduction.",
                "Core Strength: Effectively smooths out discrete step-function splits, suppressing prediction variance on volatile test days."
            ]
        },
        {
            "title": "ML Algorithm 2: Decision Tree Regressor",
            "subtitle": "Single-Tree CART Architecture with Hyperparameter Regularization",
            "bullets": [
                "Methodology: Standard Classification and Regression Tree (CART) using Mean Squared Error (MSE) criterion.",
                "Regularization: Controlled depth (max_depth=6, min_samples_split=10, min_samples_leaf=8) to mitigate overfitting.",
                "Test Performance: MAE = 14.645 kg | RMSE = 19.683 kg | R² Score = 0.4871.",
                "Core Strength: Highest interpretability; provides transparent, deterministic if-then decision paths that mess supervisors can audit."
            ]
        },
        {
            "title": "Empirical Algorithm Benchmark & Comparison",
            "subtitle": "Fair, Rigorous Comparison on Held-Out Chronological Test Set (301 Meals)",
            "bullets": [
                "Random Forest Regressor: MAE = 14.229 kg | RMSE = 19.514 kg | R² = 0.4959 | Baseline Gain: +7.68%",
                "Decision Tree Regressor: MAE = 14.645 kg | RMSE = 19.683 kg | R² = 0.4871 | Baseline Gain: +4.98%",
                "Historical Baseline: MAE = 15.412 kg | RMSE = 20.835 kg | R² = 0.4254 | Baseline Gain: 0.00%",
                "Selection Rationale: Random Forest is chosen as the primary production model due to lower error (0.416 kg MAE edge) and variance reduction; Decision Tree remains available as an interpretable secondary model."
            ]
        },
        {
            "title": "Risk Classifier & Streamlit Dashboard",
            "subtitle": "Operational Tools for Mess Supervisors",
            "bullets": [
                "Independent Classifier: Predicts high-waste risk (>35 kg) with 83.39% Accuracy, 77.56% Precision, 88.97% Recall, and 82.88% F1.",
                "Dashboard Hub: Streamlit interface featuring Theme Toggle (Dark/Light), Analytics, Live Forecasting, and Model Evaluation.",
                "Algorithm Selector: Allows users to seamlessly switch between Random Forest and Decision Tree regressors.",
                "Operational CRUD: Form-based Create, Read (multi-filter), Update (attendance/prepared/actuals), and Safe Cascading Delete."
            ]
        },
        {
            "title": "Conclusion, Verification & Viva Readiness",
            "subtitle": "Summary of Contributions & Academic Deliverables",
            "bullets": [
                "Automated Verification: Test suite with 21 unit/integration tests passing at 100% (21/21 in tests/).",
                "Academic Integrity: Genuine empirical metrics generated through systematic training and testing on identical splits.",
                "Deliverables: Python codebase, MongoDB scripts, evaluation plots, interactive dashboard, PDF report, and presentation slides.",
                "Future Scope: Integration of IoT load-cell scales and automated redistribution alerts to local charitable food networks."
            ]
        }
    ]

    total_slides = len(slides_content)
    for idx, slide_info in enumerate(slides_content):
        slide = prs.slides.add_slide(blank_layout)

        # Slide Number Badge
        badge_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.4), Inches(1.5), Inches(0.5))
        badge_tf = badge_box.text_frame
        badge_p = badge_tf.paragraphs[0]
        badge_p.text = f"Slide {idx + 1} / {total_slides}"
        badge_p.font.size = Pt(11)
        badge_p.font.color.rgb = MUTED_COLOR
        badge_p.alignment = PP_ALIGN.RIGHT

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = slide_info["title"]
        p.font.size = Pt(25)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_COLOR

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11), Inches(0.5))
        sub_tf = sub_box.text_frame
        sub_p = sub_tf.paragraphs[0]
        sub_p.text = slide_info["subtitle"]
        sub_p.font.size = Pt(13)
        sub_p.font.color.rgb = ACCENT_COLOR if idx == 0 else MUTED_COLOR

        # Content Card Box
        card_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.8))
        c_tf = card_box.text_frame
        c_tf.word_wrap = True

        for b_idx, bullet in enumerate(slide_info["bullets"]):
            bp = c_tf.add_paragraph() if b_idx > 0 else c_tf.paragraphs[0]
            bp.text = f"•  {bullet}"
            bp.font.size = Pt(14)
            bp.font.color.rgb = TEXT_COLOR
            bp.space_after = Pt(12)

    pptx_path = config.BASE_DIR / "presentation.pptx"
    prs.save(str(pptx_path))
    print(f"Presentation saved successfully to: {pptx_path}")
    return str(pptx_path)


if __name__ == "__main__":
    create_presentation()
